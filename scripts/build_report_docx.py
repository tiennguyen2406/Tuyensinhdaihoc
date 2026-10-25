"""EduRadar V2 - xuat docs/BAO_CAO_DO_AN.md thanh Word va chinh trinh bay.

Buoc 1: pypandoc (pandoc) md -> docx, co muc luc, KHONG danh so muc (tieu de da co so).
Buoc 2: python-docx hau xu ly:
  - le 2 cm; chu chinh 12 pt, gian dong 1,15
  - moi bang: style "Table Grid" (ke o), rong 100% vung chu, do rong cot ti le theo
    noi dung dai nhat cua cot; chu va code trong bang 9 pt; hang tieu de dam, nen xam, lap lai khi sang trang
  - "Phu luc C" dat trong section kho ngang

Chay: python scripts/build_report_docx.py [--md docs/BAO_CAO_DO_AN.md] [--out docs/BAO_CAO_DO_AN.docx]
Can: pip install pypandoc_binary python-docx
"""

from __future__ import annotations

import argparse
import copy
import re
import tempfile
from pathlib import Path

import pypandoc
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
BODY_PT, TABLE_PT, CODE_PT = 12, 9, 9
LANDSCAPE_HEADING = "Phụ lục C"
HEADER_FILL = "D9D9D9"
CODE_STYLES = {"Verbatim Char", "VerbatimChar"}
COL_CAP = 45  # cot mo ta dai hon nguong nay khong duoc them be rong (tranh lan cot khac)


def pandoc_docx(md: Path, out: Path) -> None:
    pypandoc.convert_file(
        str(md),
        "docx",
        # cho phep danh sach ngay sau cau dan (khong dong trong) nhu GitHub, khong bi gop vao doan
        format="markdown+lists_without_preceding_blankline-implicit_figures",
        outputfile=str(out),
        extra_args=[
            f"--resource-path={md.parent}",
        ],
    )


def add_table_grid_style(doc) -> None:
    """Reference doc cua pandoc khong co "Table Grid" -> tao style bang ke don moi canh."""
    if any(s.name == "Table Grid" for s in doc.styles):
        return
    st = OxmlElement("w:style")
    st.set(qn("w:type"), "table")
    st.set(qn("w:styleId"), "TableGrid")
    name = OxmlElement("w:name")
    name.set(qn("w:val"), "Table Grid")
    st.append(name)
    tbl_pr = OxmlElement("w:tblPr")
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        b = OxmlElement(f"w:{edge}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), "4")
        b.set(qn("w:space"), "0")
        b.set(qn("w:color"), "D9D9D9")
        borders.append(b)
    tbl_pr.append(borders)
    mar = OxmlElement("w:tblCellMar")
    for edge, w in (("left", 70), ("right", 70)):
        m = OxmlElement(f"w:{edge}")
        m.set(qn("w:w"), str(w))
        m.set(qn("w:type"), "dxa")
        mar.append(m)
    tbl_pr.append(mar)
    st.append(tbl_pr)
    doc.styles.element.append(st)


def cell_text(cell) -> str:
    """Ca chu thuong (w:t) lan cong thuc (m:t) - cot cong thuc can duoc tinh be rong."""
    return "".join(el.text or "" for el in cell._tc.iter() if el.tag in (qn("w:t"), qn("m:t")))


CHAR_TWIPS = 100  # be rong TB 1 ky tu chu 9 pt (~0,5 em; 1 pt = 20 twips) + du cho chu dam
CELL_PAD_TWIPS = 160  # le trong o trai + phai


def column_stats(table) -> list[tuple[float, int]]:
    """(trong so theo do dai noi dung, do rong toi thieu de tu dai nhat khong bi ngat)."""
    stats = []
    for j in range(len(table.columns)):
        texts = [cell_text(row.cells[j]) for row in table.rows if j < len(row.cells)]
        longest = max((len(t) for t in texts), default=1)
        word = max((len(w) for t in texts for w in t.split()), default=1)
        stats.append((max(min(longest, COL_CAP), word + 1, 3), word * CHAR_TWIPS + CELL_PAD_TWIPS))
    return stats


def fit_widths(stats: list[tuple[float, int]], total: int) -> list[int]:
    """Chia ti le theo trong so, roi nang cot hep len muc toi thieu, lay phan thieu tu cac cot du."""
    weights = [w for w, _ in stats]
    widths = [total * w / sum(weights) for w in weights]
    mins = [min(m, total / len(stats)) for _, m in stats]  # cot nao cung khong doi qua 1/n bang
    for _ in range(5):
        short = [i for i, w in enumerate(widths) if w < mins[i]]
        if not short:
            break
        need = sum(mins[i] - widths[i] for i in short)
        spare = {
            i: widths[i] - mins[i]
            for i in range(len(widths))
            if i not in short and widths[i] > mins[i]
        }
        if not spare:
            break
        for i in short:
            widths[i] = mins[i]
        tot_spare = sum(spare.values())
        for i, s in spare.items():
            widths[i] -= need * s / tot_spare
    return [int(w) for w in widths]


def set_widths(table, text_twips: int) -> None:
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    for tag in ("w:tblW", "w:tblLayout"):
        old = tbl_pr.find(qn(tag))
        if old is not None:
            tbl_pr.remove(old)
    tbl_w = OxmlElement("w:tblW")
    tbl_w.set(qn("w:w"), "5000")
    tbl_w.set(qn("w:type"), "pct")
    tbl_pr.append(tbl_w)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)

    widths = fit_widths(column_stats(table), text_twips)
    if [cell_text(c) for c in table.rows[0].cells] == [
        "STT",
        "ID",
        "Trường",
        "Năm",
        "Phương thức",
        "Chương trình",
        "Tổ hợp",
        "Điểm",
        "Thang",
        "Mã ngành",
        "Đúng/Sai",
    ]:
        # Phu luc C: ma phuong thuc can du rong; danh nhieu cho ten va nhom to hop.
        ratios = [0.8, 1.1, 1.4, 1.1, 2.3, 7.8, 5.2, 1.3, 1.3, 1.7, 1.7]
        widths = [int(text_twips * r / sum(ratios)) for r in ratios]
    grid = tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(w))
    for row in table.rows:
        for cell, w in zip(row.cells, widths):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.insert(0, tc_w)
            tc_w.set(qn("w:w"), str(w))
            tc_w.set(qn("w:type"), "dxa")


def size_math_runs(cell, pt: int) -> None:
    for mr in cell._tc.iter(qn("m:r")):
        rpr = mr.find(qn("w:rPr"))
        if rpr is None:
            rpr = OxmlElement("w:rPr")
            mrpr = mr.find(qn("m:rPr"))
            mr.insert(0 if mrpr is None else list(mr).index(mrpr) + 1, rpr)
        sz = rpr.find(qn("w:sz"))
        if sz is None:
            sz = OxmlElement("w:sz")
            rpr.append(sz)
        sz.set(qn("w:val"), str(pt * 2))


def format_table(table, text_twips: int) -> None:
    table.style = "Table Grid"
    set_widths(table, text_twips)
    for i, row in enumerate(table.rows):
        if i == 0:  # hang tieu de: lap lai khi sang trang
            tr_pr = row._tr.get_or_add_trPr()
            if tr_pr.find(qn("w:tblHeader")) is None:
                tr_pr.append(OxmlElement("w:tblHeader"))
        tr_pr = row._tr.get_or_add_trPr()
        tr_pr.append(OxmlElement("w:cantSplit"))
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if i == 0:
                shd = OxmlElement("w:shd")
                shd.set(qn("w:val"), "clear")
                shd.set(qn("w:color"), "auto")
                shd.set(qn("w:fill"), HEADER_FILL)
                cell._tc.get_or_add_tcPr().append(shd)
            for p in cell.paragraphs:
                pf = p.paragraph_format
                pf.space_before = pf.space_after = Pt(1)
                pf.line_spacing = 1.0
                for r in p.runs:
                    is_code = r.style is not None and r.style.name in CODE_STYLES
                    r.font.size = Pt(CODE_PT if is_code else TABLE_PT)
                    r.font.name = "Times New Roman"
                    if i == 0:
                        r.font.bold = True
            size_math_runs(cell, TABLE_PT)
            # paragraph.runs khong bao gom run nam trong hyperlink cua Pandoc.
            for run in cell._tc.iter(qn("w:r")):
                rpr = run.find(qn("w:rPr"))
                if rpr is None:
                    rpr = OxmlElement("w:rPr")
                    run.insert(0, rpr)
                for tag in ("w:sz", "w:szCs"):
                    size = rpr.find(qn(tag))
                    if size is None:
                        size = OxmlElement(tag)
                        rpr.append(size)
                    size.set(qn("w:val"), str(TABLE_PT * 2))


def text_width_twips(section) -> int:
    return int((section.page_width - section.left_margin - section.right_margin) / 635)


def landscape_from(doc, heading: str) -> None:
    """Chen ngat section truoc heading; section cuoi (tu heading den het) thanh kho ngang."""
    target = next((p for p in doc.paragraphs if p.text.strip().startswith(heading)), None)
    if target is None:
        raise SystemExit(f"Khong tim thay tieu de '{heading}'")
    body_sect = doc.element.body.sectPr
    brk = OxmlElement("w:p")
    ppr = OxmlElement("w:pPr")
    ppr.append(copy.deepcopy(body_sect))  # section truoc giu kho doc
    brk.append(ppr)
    target._p.addprevious(brk)
    last = doc.sections[-1]
    w, h = last.page_width, last.page_height
    last.orientation = WD_ORIENT.LANDSCAPE
    last.page_width, last.page_height = max(w, h), min(w, h)


def postprocess(path: Path) -> dict:
    doc = Document(str(path))
    for sec in doc.sections:  # pandoc khong ghi kho giay -> dat ro A4 truoc khi tach section
        sec.page_width, sec.page_height = Cm(21), Cm(29.7)
        sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Cm(2)
    for name in ("Normal", "Body Text", "First Paragraph", "Compact"):
        st = next((s for s in doc.styles if s.name == name), None)
        if st is not None:
            st.font.size = Pt(BODY_PT)
            st.font.name = "Times New Roman"
            st.paragraph_format.line_spacing = 1.15
            st.paragraph_format.space_after = Pt(6)
    for style in doc.styles:
        if style.type == 1:
            style.font.color.rgb = RGBColor(0, 0, 0)
            if style.name.startswith("Heading"):
                style.font.name = "Times New Roman"
                style.font.bold = True
                style.font.size = Pt(
                    16 if style.name == "Heading 1" else 14 if style.name == "Heading 2" else 12
                )
                style.paragraph_format.keep_with_next = True
                if style.name == "Heading 1":
                    style.paragraph_format.page_break_before = True
    for p in doc.paragraphs:
        if p.style.name in ("Title", "Subtitle"):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(55 if p.style.name == "Title" else 18)
            p.paragraph_format.space_after = Pt(25)
            for r in p.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(22 if p.style.name == "Title" else 14)
                r.font.bold = True
                r.font.color.rgb = RGBColor(0, 0, 0)
        if p.text.strip() in ("[Họ tên — MSSV]", "[Lớp]", "[Giảng viên]", "[Ngày nộp]"):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(18)
        if p.text.strip() == "[[TOC]]":
            title = p.insert_paragraph_before("Mục lục")
            title.paragraph_format.page_break_before = True
            title.runs[0].font.size = Pt(16)
            title.runs[0].bold = True
            p.clear()
            for kind in ("begin", "separate", "end"):
                run = p.add_run()
                if kind == "separate":
                    instr = OxmlElement("w:instrText")
                    instr.set(qn("xml:space"), "preserve")
                    instr.text = ' TOC \\o "1-2" \\h \\z \\u '
                    run._r.append(instr)
                fld = OxmlElement("w:fldChar")
                fld.set(qn("w:fldCharType"), kind)
                run._r.append(fld)
        if p._p.xpath(".//w:drawing"):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True
        if re.match(r"^Hình \d+\.\d+\.", p.text):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_together = True
            for r in p.runs:
                r.font.size = Pt(10)
        if p.text.startswith("Nguồn:"):
            for r in p.runs:
                r.font.size = Pt(10)
        if p.text.startswith("Bảng "):
            p.paragraph_format.keep_with_next = True
    for shape in doc.inline_shapes:
        # Khong chi thu nho chieu rong: gioi han ca chieu cao de caption cung trang.
        factor = min(1, Cm(17) / shape.width, Cm(16) / shape.height)
        shape.width, shape.height = int(shape.width * factor), int(shape.height * factor)
    settings = doc.settings.element
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    settings.append(update)
    add_table_grid_style(doc)
    landscape_from(doc, LANDSCAPE_HEADING)

    # bang nam sau heading Phu luc C -> do rong theo kho ngang
    body = doc.element.body
    heading_p = next(p for p in doc.paragraphs if p.text.strip().startswith(LANDSCAPE_HEADING))
    order = list(body)
    land_start = order.index(heading_p._p)
    portrait_tw, land_tw = text_width_twips(doc.sections[0]), text_width_twips(doc.sections[-1])
    n_land = 0
    for t in doc.tables:
        landscape = order.index(t._tbl) > land_start
        n_land += landscape
        format_table(t, land_tw if landscape else portrait_tw)
    # Cac section lien ket cung footer; chi them field mot lan.
    p = doc.sections[0].footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    p._p.append(field)
    doc.save(str(path))
    return {"tables": len(doc.tables), "landscape_tables": n_land, "sections": len(doc.sections)}


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--md", type=Path, default=ROOT / "docs" / "BAO_CAO_DO_AN.md")
    ap.add_argument("--out", type=Path, default=ROOT / "docs" / "BAO_CAO_DO_AN.docx")
    args = ap.parse_args()
    with tempfile.TemporaryDirectory() as tmp:
        raw = Path(tmp) / "raw.docx"
        pandoc_docx(args.md, raw)
        info = postprocess(raw)
        args.out.write_bytes(raw.read_bytes())
    print(f"Exported {args.out.relative_to(ROOT)}: {info}")
    # kiem khong con so muc lap kieu "1 1." hay "2.1 2.1"
    heads = [
        p.text for p in Document(str(args.out)).paragraphs if p.style.name.startswith("Heading")
    ]
    doubled = [h for h in heads if re.match(r"^\d+(\.\d+)*\.?\s+\d", h)]
    print("Repeated heading numbers:", len(doubled))


if __name__ == "__main__":
    main()
